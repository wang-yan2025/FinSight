"""核验报告结构化与导出（Markdown / DOCX）。"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from finsight.audit import audit_event
from finsight.verify.rules import STATUS_CN, VerificationReportData

VERIFY_DISCLAIMER = (
    "本核验报告由 FinSight 智研投基于东方财富结构化财务数据与确定性计算引擎独立重算生成："
    "pass 仅代表该条声明经程序比对（数值、指标、单位、期间、公司、引用关系）一致；"
    "「通过数量」仅指成功提取并完成核验的声明，报告其余内容未经提取或核验，"
    "不等于整份报告完全正确；声明提取基于保守正则规则，识别覆盖有限，"
    "无法识别的内容不会被自动判定正确；"
    "参考数据与备忘录数据同源（东方财富结构化财务数据），本核验构成计算路径独立比对，"
    "不构成跨数据源交叉验证；"
    "insufficient_evidence 表示缺少证据支持，不等于声明虚假；unable_to_verify 表示核验条件不足；"
    "needs_review 表示需要人工判断。核验结果不构成对报告整体质量或投资价值的认定。"
)


def _status_badge(status: str) -> str:
    """状态标记：错误与待复核事项视觉区分。"""
    marks = {"pass": "✔", "fail": "✘", "insufficient_evidence": "○",
             "unable_to_verify": "○", "needs_review": "⚠"}
    return f"{marks.get(status, '·')}【{STATUS_CN.get(status, status)}】"


def _result_lines(r) -> list[str]:
    lines = [f"- {_status_badge(r.status)} `{r.cid}`（{r.location}）",
             f"  - 原文：{r.claim_text}"]
    if r.claimed_display:
        lines.append(f"  - 声称值：{r.claimed_display}")
    if r.reference_display:
        lines.append(f"  - 参考值：{r.reference_display}（{r.ref_period}，{r.ref_basis}）")
    if r.difference:
        lines.append(f"  - 差异：{r.difference}")
    if r.diff_reason:
        lines.append(f"  - 原因：{r.diff_reason}")
    if r.basis_detail:
        lines.append(f"  - 核验依据：{r.basis_detail}；来源：{r.ref_source}")
    if r.suggestion:
        lines.append(f"  - 建议修订：{r.suggestion}")
    for lim in r.limitations:
        lines.append(f"  - 限制：{lim}")
    return lines


def report_to_markdown(report: VerificationReportData) -> str:
    counts = report.counts()
    lines = [
        f"# 研究报告核验报告——{report.code}",
        "",
        f"- 核验对象：{report.target_desc}",
        f"- 核验时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"- 声明总数：{report.claim_count}（仅指成功提取并进入核验的声明，"
        "不代表报告全部内容已核验）；"
        f"通过 {counts['pass']}／不通过 {counts['fail']}／证据不足 {counts['insufficient_evidence']}"
        f"／无法核验 {counts['unable_to_verify']}／需人工复核 {counts['needs_review']}",
        f"- 参考数据来源：东方财富三大报表（经 akshare，本地缓存快照）+ 既有指标引擎独立重算"
        "（与备忘录数据同源，属计算路径独立比对，非跨数据源交叉验证）",
        f"- 快照说明：{report.snapshot_note or '未提供报告生成时证据登记'}",
        "",
        "## 各状态结果",
        "",
    ]
    for status in ("fail", "insufficient_evidence", "unable_to_verify", "needs_review", "pass"):
        group = [r for r in report.results if r.status == status]
        if not group:
            continue
        lines.append(f"### {STATUS_CN[status]}（{len(group)}）")
        lines.append("")
        for r in group:
            lines.extend(_result_lines(r))
        lines.append("")
    if report.source_findings:
        lines.append("## 来源描述检查")
        lines.append("")
        lines.extend(f"- {f}" for f in report.source_findings)
        lines.append("")
    lines.append("## 人工复核清单")
    lines.append("")
    needs = [r for r in report.results if r.status in ("needs_review", "unable_to_verify",
                                                       "insufficient_evidence")]
    if needs:
        lines.extend(f"- `{r.cid}`：{r.claim_text[:60]}…（{STATUS_CN[r.status]}）" for r in needs)
    else:
        lines.append("- 无需人工复核项")
    lines.append("")
    lines.append("---")
    lines.append(f"**声明**：{VERIFY_DISCLAIMER}")
    audit_event("export_report", format="markdown", target=report.code,
                results=len(report.results), chars=len("\n".join(lines)))
    return "\n".join(lines)


def report_to_docx_bytes(report: VerificationReportData) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    doc.add_heading(f"研究报告核验报告——{report.code}", level=0)
    counts = report.counts()
    doc.add_paragraph(
        f"核验对象：{report.target_desc}\n"
        f"核验时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"声明总数：{report.claim_count}｜通过 {counts['pass']}｜不通过 {counts['fail']}"
        f"｜证据不足 {counts['insufficient_evidence']}｜无法核验 {counts['unable_to_verify']}"
        f"｜需人工复核 {counts['needs_review']}\n"
        f"快照说明：{report.snapshot_note or '未提供报告生成时证据登记'}")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for i, head in enumerate(["ID", "状态", "原文（节选）", "声称 vs 参考", "建议"]):
        table.rows[0].cells[i].text = head
    for r in report.results:
        row = table.add_row().cells
        row[0].text = r.cid
        row[1].text = _status_badge(r.status)
        row[2].text = r.claim_text[:80]
        row[3].text = (f"{r.claimed_display} vs {r.reference_display}"
                       + (f"（{r.difference}）" if r.difference else ""))
        row[4].text = r.suggestion or ("—".join(r.limitations[:1]) if r.limitations else "—")
    doc.add_page_break()

    if report.source_findings:
        doc.add_heading("来源描述检查", level=1)
        for f in report.source_findings:
            doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("人工复核清单", level=1)
    needs = [r for r in report.results if r.status in ("needs_review", "unable_to_verify",
                                                       "insufficient_evidence")]
    if needs:
        for r in needs:
            doc.add_paragraph(f"{r.cid}：{r.claim_text[:80]}（{STATUS_CN[r.status]}）",
                              style="List Bullet")
    else:
        doc.add_paragraph("无需人工复核项")

    doc.add_heading("声明", level=1)
    doc.add_paragraph(VERIFY_DISCLAIMER)

    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    audit_event("export_report", format="docx", target=report.code,
                results=len(report.results), bytes=len(data))
    return data
