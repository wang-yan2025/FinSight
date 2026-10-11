"""LLM 层测试：mock 请求，覆盖配置/成功/失败/校验；导出测试。"""
import io
import json

import pandas as pd
import pytest
from docx import Document

from finsight.finance.field_mapping import REPORT_DATE_FIELD, standardize_statement
from finsight.memo import llm as llm_mod
from finsight.memo.evidence import build_evidence_book
from finsight.memo.llm import (
    LLMConfig,
    LLMError,
    apply_llm_sections,
    build_number_allowlist,
    generate_llm_sections,
    validate_llm_sections,
)
from finsight.memo.render import memo_to_docx_bytes, memo_to_markdown
from finsight.memo.skeleton import build_memo_skeleton


def qd(year, quarter):
    m, d = {1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}[quarter]
    return pd.Timestamp(year, m, d)


def cum(singles):
    total, out = 0, []
    for s in singles:
        total += s
        out.append(total)
    return out


def build_memo():
    dates = [qd(2023, q) for q in (1, 2, 3, 4)] + [qd(2024, q) for q in (1, 2, 3, 4)]
    inc = standardize_statement("income", pd.DataFrame({
        REPORT_DATE_FIELD: [d.strftime("%Y-%m-%d %H:%M:%S") for d in dates],
        "OPERATE_INCOME": cum([100] * 4) + cum([120, 110, 90, 95]),
        "OPERATE_COST": cum([60] * 4) + cum([70, 66, 54, 57]),
        "NETPROFIT": cum([50] * 4) + cum([10, -5, -10, -15]),
        "PARENT_NETPROFIT": cum([45] * 4) + cum([9, -4, -9, -13]),
    }))
    book = build_evidence_book("T", inc, None, None)
    return build_memo_skeleton("T", book)


VALID_EIDS = {"E-REV-2024Q4", "E-NPYOY-2024Q4"}


# ---------------- 配置 ----------------

def test_config_from_env(monkeypatch):
    monkeypatch.delenv(llm_mod.ENV_KEY, raising=False)
    monkeypatch.delenv(llm_mod.ENV_BASE, raising=False)
    monkeypatch.delenv(llm_mod.ENV_MODEL, raising=False)
    cfg = LLMConfig.from_env()
    assert not cfg.configured()
    monkeypatch.setenv(llm_mod.ENV_KEY, "sk-test")
    monkeypatch.setenv(llm_mod.ENV_BASE, "https://api.example.com/v1")
    monkeypatch.setenv(llm_mod.ENV_MODEL, "test-model")
    cfg = LLMConfig.from_env()
    assert cfg.configured()


def test_call_llm_unconfigured_raises():
    with pytest.raises(LLMError):
        llm_mod.call_llm(LLMConfig(), "s", "u")


# ---------------- mock 调用 ----------------

class _FakeResponse:
    def __init__(self, status_code=200, payload=None):
        self.status_code = status_code
        self._payload = payload or {}

    def json(self):
        return self._payload


def _ok_payload(text):
    return {"choices": [{"message": {"content": text}}]}


def test_call_llm_success(monkeypatch):
    captured = {}

    def fake_post(url, json=None, headers=None, timeout=None):
        captured.update(url=url, body=json, headers=headers)
        return _FakeResponse(payload=_ok_payload('{"summary": "ok"}'))

    monkeypatch.setattr(llm_mod.requests, "post", fake_post)
    cfg = LLMConfig(api_key="sk-test", base_url="https://api.example.com/v1", model="m1")
    out = llm_mod.call_llm(cfg, "sys", "user")
    assert out == '{"summary": "ok"}'
    # 密钥只出现在请求头，不出现在请求体/URL
    assert captured["headers"]["Authorization"] == "Bearer sk-test"
    assert "sk-test" not in json.dumps(captured["body"]) and "sk-test" not in captured["url"]


def test_call_llm_http_error(monkeypatch):
    monkeypatch.setattr(llm_mod.requests, "post",
                        lambda *a, **k: _FakeResponse(status_code=500))
    cfg = LLMConfig(api_key="k", base_url="https://x/v1", model="m")
    with pytest.raises(LLMError, match="HTTP 500"):
        llm_mod.call_llm(cfg, "s", "u")


def test_call_llm_network_error(monkeypatch):
    import requests as _r

    def boom(*a, **k):
        raise _r.ConnectionError("refused")

    monkeypatch.setattr(llm_mod.requests, "post", boom)
    cfg = LLMConfig(api_key="k", base_url="https://x/v1", model="m")
    with pytest.raises(LLMError, match="网络或接口请求异常"):
        llm_mod.call_llm(cfg, "s", "u")


def test_call_llm_bad_structure(monkeypatch):
    monkeypatch.setattr(llm_mod.requests, "post",
                        lambda *a, **k: _FakeResponse(payload={"foo": 1}))
    cfg = LLMConfig(api_key="k", base_url="https://x/v1", model="m")
    with pytest.raises(LLMError, match="响应结构异常"):
        llm_mod.call_llm(cfg, "s", "u")


def test_generate_llm_sections_end_to_end(monkeypatch):
    memo = build_memo()
    valid_ids = set(memo.evidence.ids())
    some_id = "E-REV-2024Q4"

    def fake_post(url, json=None, headers=None, timeout=None):
        content = json_lib.dumps({
            "summary": f"收入 415 元，正增长 [{some_id}]",
            "company": "公司业务材料未接入，待补充。",
            "thesis": f"利润下滑是主要反面因素 [{some_id}]",
            "scenarios": "定性情景：乐观/基准/悲观（无估值数字）。",
        }, ensure_ascii=False)
        return _FakeResponse(payload=_ok_payload(content))

    import json as json_lib
    monkeypatch.setattr(llm_mod.requests, "post", fake_post)
    cfg = LLMConfig(api_key="k", base_url="https://x/v1", model="m")
    sections, issues = generate_llm_sections(cfg, memo, focus_question="")
    assert set(sections) == {"summary", "company", "thesis", "scenarios"}
    assert not issues
    apply_llm_sections(memo, sections, issues)
    assert memo.llm_status == "已生成"
    assert memo.section("company").llm_text


# ---------------- 输出校验 ----------------

def _allowlist():
    memo = build_memo()
    return build_number_allowlist(memo)


def test_validate_ok_output():
    data = {"summary": f"营业收入正增长 [E-REV-2024Q4]，净利下滑明显 [E-NPYOY-2024Q4]。",
            "company": "公司概况待补充。", "thesis": "事实与推断区分。", "scenarios": "定性情景。"}
    sections, issues = validate_llm_sections(data, VALID_EIDS, _allowlist())
    assert set(sections) == set(data) and not issues


def test_validate_invalid_eid_marked():
    data = {"summary": "收入高增 [E-FAKE-2099Q9]。", "company": "x", "thesis": "y", "scenarios": "z"}
    sections, issues = validate_llm_sections(data, VALID_EIDS, _allowlist())
    assert "[待核实引用]" in sections["summary"]
    assert any("无效证据 ID" in i for i in issues)
    assert "E-FAKE-2099Q9" not in sections["summary"]


def test_validate_unknown_numbers_flagged():
    data = {"summary": "预计营收将达 999.5 亿元。", "company": "x", "thesis": "y", "scenarios": "z"}
    sections, issues = validate_llm_sections(data, VALID_EIDS, _allowlist())
    assert "待核实" in sections["summary"]
    assert any("未核验数字" in i for i in issues)


def test_validate_allowed_numbers_pass():
    memo = build_memo()
    allow = build_number_allowlist(memo)
    data = {"summary": f"营业收入累计 415 元 [E-REV-2024Q4]，同比约 +3.75%。",
            "company": "x", "thesis": "y", "scenarios": "z"}
    # 白名单包含证据数值的多种格式；415/3.75 均来自证据
    sections, issues = validate_llm_sections(data, VALID_EIDS, allow)
    assert sections["summary"].endswith("]") or "待核实" not in sections["summary"]
    assert not [i for i in issues if "未核验数字" in i]


def test_validate_missing_sections_and_unknown_keys():
    data = {"summary": "s", "extra_key": "被忽略"}
    sections, issues = validate_llm_sections(data, VALID_EIDS, _allowlist())
    assert list(sections) == ["summary"]
    assert any("缺失" in i for i in issues)
    assert any("未知输出键" in i for i in issues)


def test_parse_json_with_code_fence():
    text = "```json\n{\"summary\": \"ok\"}\n```"
    assert llm_mod._parse_json_text(text)["summary"] == "ok"


def test_parse_json_invalid_raises():
    with pytest.raises(LLMError, match="有效 JSON"):
        llm_mod._parse_json_text("not json at all")


# ---------------- 导出 ----------------

def test_markdown_export_complete():
    memo = build_memo()
    md = memo_to_markdown(memo)
    for key in ("研究摘要", "财务表现", "关键风险", "证据清单", "免责声明"):
        assert key in md
    assert "E-REV-2024Q4" in md
    assert "待核实" in md or "待补充" in md
    assert "2024Q4" in md


def test_docx_export_complete():
    memo = build_memo()
    data = memo_to_docx_bytes(memo)
    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("研究摘要" in h for h in headings)
    assert any("证据清单" in h for h in headings)
    # 证据附录表格：表头 + 证据行
    assert doc.tables, "应有证据附录表格"
    header_cells = [c.text for c in doc.tables[0].rows[0].cells]
    assert "证据 ID" in header_cells
    assert len(doc.tables[0].rows) > 3
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "免责声明" in body_text and "E-REV-2024Q4" in body_text


def test_docx_with_llm_sections():
    memo = build_memo()
    sections = {"summary": "LLM 摘要内容。", "company": "公司概况（LLM）。",
                "thesis": "投资逻辑（LLM）。", "scenarios": "情景（LLM）。"}
    apply_llm_sections(memo, sections, [])
    doc = Document(io.BytesIO(memo_to_docx_bytes(memo)))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "LLM 摘要内容" in body and "尚未经独立核验" in body
