"""备忘录渲染与导出：Markdown 与 Word（python-docx，中文正常显示）。"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from finsight.audit import audit_event
from finsight.memo.llm import LLM_PROVENANCE_NOTE
from finsight.memo.skeleton import DISCLAIMER, InvestmentMemo

_LLM_BANNER = f"【{LLM_PROVENANCE_NOTE}】"


def memo_to_markdown(memo: InvestmentMemo) -> str:
    """完整备忘录 → Markdown（含证据引用、快照、待核实与免责声明）。"""
    lines = [
        f"# 买方投资备忘录（研究辅助稿）——{memo.code}",
        "",
        f"- 生成时间：{memo.generated_at}｜目标报告期：{memo.period}",
        f"- LLM 状态：{memo.llm_status}" +
        (f"（{len(memo.llm_issues)} 条校验提示）" if memo.llm_issues else ""),
        "",
    ]
    for sec in memo.sections:
        lines.append(f"## {sec.title}")
        lines.append("")
        if sec.llm_text:
            lines.append(f"> {_LLM_BANNER}")
            lines.append(f"> {sec.llm_text}")
            lines.append("")
        for fact in sec.facts:
            lines.append(fact)
        if sec.key == "appendix" and memo.evidence is not None:
            lines.append("")
            lines.append("| 证据 ID | 内容 | 值 | 单位 | 报告期 | 来源字段/依据 |")
            lines.append("|---|---|---|---|---|---|")
            for eid in memo.evidence.ids():
                e = memo.evidence.get(eid)
                value = "—" if e.value is None else f"{e.value:,.4g}"
                lines.append(f"| {e.eid} | {e.label} | {value} | {e.unit} | "
                             f"{e.report_date} | {e.derivation} |")
        lines.append("")
    lines.append("---")
    lines.append(f"**免责声明**：{DISCLAIMER}")
    lines.append("")
    audit_event("export_memo", format="markdown", target=memo.code,
                period=memo.period, sections=len(memo.sections),
                chars=len("\n".join(lines)))
    return "\n".join(lines)


def _set_cjk_font(doc: Document) -> None:
    """中文正文用宋体、西文 Calibri（Word 中正常渲染中文）。"""
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def memo_to_docx_bytes(memo: InvestmentMemo) -> bytes:
    """完整备忘录 → DOCX 字节流（标题层级、证据表、免责声明）。"""
    doc = Document()
    _set_cjk_font(doc)

    title = doc.add_heading(f"买方投资备忘录（研究辅助稿）——{memo.code}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        f"生成时间：{memo.generated_at}｜目标报告期：{memo.period}｜LLM 状态：{memo.llm_status}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for sec in memo.sections:
        doc.add_heading(sec.title, level=1)
        if sec.llm_text:
            doc.add_paragraph(_LLM_BANNER)
            doc.add_paragraph(sec.llm_text)
        for fact in sec.facts:
            text = fact.strip("*").strip()
            if fact.startswith("**"):  # 组名 → 二级标题
                doc.add_heading(text, level=2)
            else:
                doc.add_paragraph(text, style="List Bullet" if fact.startswith("- ") else None)
        if sec.key == "appendix" and memo.evidence is not None:
            table = doc.add_table(rows=1, cols=6)
            table.style = "Light Grid Accent 1"
            for i, head in enumerate(["证据 ID", "内容", "值", "单位", "报告期", "来源字段/依据"]):
                table.rows[0].cells[i].text = head
            for eid in memo.evidence.ids():
                e = memo.evidence.get(eid)
                row = table.add_row().cells
                value = "—" if e.value is None else f"{e.value:,.4g}"
                row[0].text, row[1].text, row[2].text = e.eid, e.label, value
                row[3].text, row[4].text, row[5].text = e.unit, e.report_date, e.derivation

    doc.add_heading("免责声明", level=1)
    doc.add_paragraph(DISCLAIMER)

    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    audit_event("export_memo", format="docx", target=memo.code,
                period=memo.period, sections=len(memo.sections), bytes=len(data))
    return data


# ---------------- PDF 导出（reportlab + 系统中文字体） ----------------

def _pdf_cjk_font():
    """注册系统 CJK 字体；不可用时返回 None（调用方降级为拉丁字体）。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for name, path in (("FinSightHei", "C:/Windows/Fonts/msyh.ttc"),
                       ("FinSightHei", "C:/Windows/Fonts/simhei.ttf"),
                       ("FinSightHei", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")):
        try:
            pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception:
            continue
    return None


def memo_to_pdf_bytes(memo: InvestmentMemo) -> bytes:
    """完整备忘录 → PDF 字节流（reportlab，中文标题/段落/证据表）。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    font = _pdf_cjk_font() or "Helvetica"
    ink = colors.HexColor("#1D2939")
    muted = colors.HexColor("#667085")
    line = colors.HexColor("#E4E9F2")
    s_title = ParagraphStyle("t", fontName=font, fontSize=16, leading=22, textColor=ink,
                             alignment=1, spaceAfter=4 * mm)
    s_meta = ParagraphStyle("m", fontName=font, fontSize=8.5, leading=12, textColor=muted,
                            alignment=1, spaceAfter=6 * mm)
    s_h1 = ParagraphStyle("h1", fontName=font, fontSize=12, leading=17, textColor=ink,
                          spaceBefore=5 * mm, spaceAfter=2.5 * mm)
    s_body = ParagraphStyle("b", fontName=font, fontSize=9.5, leading=14.5, textColor=ink)
    s_llm = ParagraphStyle("l", fontName=font, fontSize=9, leading=13.5,
                           textColor=colors.HexColor("#51607A"))

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            title=f"FinSight 备忘录 {memo.code}")
    story = [Paragraph(esc(f"买方投资备忘录（研究辅助稿）——{memo.code}"), s_title),
             Paragraph(esc(f"生成时间：{memo.generated_at}｜目标报告期：{memo.period}"
                           f"｜LLM 状态：{memo.llm_status}"), s_meta)]
    for sec in memo.sections:
        story.append(Paragraph(esc(sec.title), s_h1))
        if sec.llm_text:
            story.append(Paragraph(esc(_LLM_BANNER), s_llm))
            story.append(Paragraph(esc(sec.llm_text), s_body))
            story.append(Spacer(1, 2 * mm))
        for fact in sec.facts:
            story.append(Paragraph(esc(fact.strip("*").strip()), s_body))
        if sec.key == "appendix" and memo.evidence is not None:
            data = [["证据 ID", "内容", "值", "单位", "报告期"]]
            for eid in memo.evidence.ids():
                e = memo.evidence.get(eid)
                data.append([e.eid, e.label[:28],
                             "—" if e.value is None else f"{e.value:,.4g}",
                             e.unit, e.report_date])
            tbl = Table(data, colWidths=[30 * mm, 62 * mm, 30 * mm, 12 * mm, 22 * mm])
            tbl.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F8")),
                ("GRID", (0, 0), (-1, -1), 0.4, line),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story.append(tbl)
    story.append(Paragraph("免责声明", s_h1))
    story.append(Paragraph(esc(DISCLAIMER), s_body))
    doc.build(story)
    pdf = buf.getvalue()
    audit_event("export_memo", format="pdf", target=memo.code,
                period=memo.period, sections=len(memo.sections), bytes=len(pdf))
    return pdf
