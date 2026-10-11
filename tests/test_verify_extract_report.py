"""提取与声明解析测试 + 核验报告导出测试 + 上传管线端到端（全离线）。"""
import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from finsight.finance.field_mapping import REPORT_DATE_FIELD, standardize_statement
from finsight.verify.claims import extract_claims
from finsight.verify.extract import (
    DocumentError,
    extract_docx,
    extract_fragments,
    extract_markdown,
    validate_upload,
)
from finsight.verify.reference import build_reference_snapshot
from finsight.verify.report import report_to_docx_bytes, report_to_markdown
from finsight.verify.rules import verify_claim_set

sys.path.insert(0, str(Path(__file__).parent))
from test_verify_rules import make_snap  # 复用固定手算数据集


# ---------------- 文件校验与提取 ----------------

def test_validate_upload_rejects_bad_inputs():
    with pytest.raises(DocumentError):
        validate_upload("report.pdf", b"x")
    with pytest.raises(DocumentError):
        validate_upload("a.md", b"")
    with pytest.raises(DocumentError):
        validate_upload("a.docx", b"x" * (10 * 1024 * 1024 + 1))


def test_extract_markdown_locations():
    frags = extract_markdown("# 标题一\n\n2024年营业收入为415亿元。\n| 指标 | 值 |\n| 6% | x |")
    assert any("标题" in f.location for f in frags)
    assert any("表格" in f.location for f in frags)
    assert frags[0].section == "标题一"


def test_extract_docx_with_table():
    doc = Document()
    doc.add_heading("测试报告", level=1)
    doc.add_paragraph("2024年营业收入为415亿元。")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "2024年净利润为-20亿元"
    buf = io.BytesIO()
    doc.save(buf)
    frags = extract_docx(buf.getvalue())
    assert any("表格 1 行 1 列 2" == f.location for f in frags)
    assert any(f.section == "测试报告" for f in frags)


def test_extract_fragments_corrupted_docx():
    with pytest.raises(DocumentError, match="无法解析"):
        extract_fragments("bad.docx", b"this is not a docx file")
    with pytest.raises(DocumentError):
        extract_fragments("bad.txt", b"hello")


# ---------------- 声明解析 ----------------

def test_claim_extraction_amount_and_period():
    frags = extract_markdown("2024年营业收入为415亿元。")
    claims = extract_claims(frags).claims
    assert len(claims) == 1
    c = claims[0]
    assert c.metric == "revenue" and c.period == "2024Q4" and c.period_basis == "cumulative"
    assert c.value == pytest.approx(415e8) and c.unit == "元" and c.decimals == 0


def test_claim_extraction_yoy_percent():
    frags = extract_markdown("2024年营业收入同比增长3.75%。")
    c = extract_claims(frags).claims[0]
    assert c.metric == "revenue_yoy" and c.value == pytest.approx(0.0375) and c.unit == "%"


def test_claim_extraction_period_variants():
    cases = {
        "2024年二季度营业收入为110亿元。": ("2024Q2", "quarter"),
        "2024Q4营业收入为415亿元。": ("2024Q4", "quarter_report"),
        "2024年上半年营业收入为230亿元。": ("2024Q2", "cumulative"),
        "2024年前三季度营业收入为320亿元。": ("2024Q3", "cumulative"),
    }
    for text, (period, basis) in cases.items():
        claims = extract_claims(extract_markdown(text)).claims
        assert claims[0].period == period and claims[0].period_basis == basis, text


def test_claim_extraction_ratio_and_pp():
    ratio = extract_claims(extract_markdown("2024年毛利率为40.5%。")).claims[0]
    assert ratio.metric == "gross_margin" and ratio.unit == "%"
    pp = extract_claims(extract_markdown("2024年营业收入同比增长3.75个百分点。")).claims[0]
    assert pp.unit == "百分点" and pp.value == pytest.approx(3.75)


def test_claim_extraction_status_word():
    c = extract_claims(extract_markdown("2024年净利润由盈转亏。")).claims[0]
    assert c.kind == "status_word" and c.metric == "net_profit_yoy" and c.period == "2024Q4"


def test_claim_extraction_qualitative_keyword():
    c = extract_claims(extract_markdown("公司拥有行业领先的竞争壁垒。")).claims[0]
    assert c.kind == "qualitative" and c.evidence_ref is None


def test_claim_extraction_evidence_ref():
    text = "营业收入（2024Q4，累计）415 元 [E-REV-2024Q4]。"
    c = extract_claims(extract_markdown(text)).claims[0]
    assert c.evidence_ref == "E-REV-2024Q4"


def test_claim_extraction_no_false_claims():
    # 无指标词或无数值 → 不生成声明（保守）
    assert extract_claims(extract_markdown("公司经营稳健，前景良好。")).claims == []
    assert extract_claims(extract_markdown("营业收入表现优秀。")).claims == []


# ---------------- 上传管线端到端（入口 B） ----------------

def test_uploaded_report_pipeline_end_to_end():
    """正确/错误混合报告 → 提取 → 独立核验：预期状态逐条手算。"""
    report_md = "\n".join([
        "# 某公司研究报告",
        "2024年营业收入为415亿元，净利润由盈转亏。",
        "2024年净利润为-20亿元。",
        "2024年二季度营业收入为230亿元。",
        "2024年营业收入同比增长3.75个百分点。",
        "2024年毛利率为40.5%。",
        "公司拥有行业领先的竞争壁垒。",
        "2021年营业收入为300亿元。",
    ])
    frags = extract_fragments("report.md", report_md.encode("utf-8"))
    snap = make_snap()
    claim_set = extract_claims(frags)
    report = verify_claim_set(claim_set, snap, fragments_text=report_md, target_desc="测试报告")

    by_metric = {(c.metric, c.kind, str(c.value)): c.cid for c in claim_set.claims}
    result_by_cid = {r.cid: r for r in report.results}
    status = lambda c: result_by_cid[by_metric[c]].status

    assert status(("revenue", "numeric", str(415e8))) == "pass"          # 正确数字
    assert status(("net_profit_yoy", "status_word", "None")) == "pass"  # 由盈转亏正确
    assert status(("net_profit", "numeric", str(-20e8))) == "pass"      # −20 亿正确
    assert status(("revenue", "numeric", str(230e8))) == "fail"         # 累计误作单季
    assert status(("revenue_yoy", "numeric", "3.75")) == "fail"         # pp 概念混淆
    assert status(("gross_margin", "numeric", str(0.405))) == "pass"    # 40.5% 容差内
    qual = [c for c in claim_set.claims if c.kind == "qualitative"]
    assert qual and result_by_cid[qual[0].cid].status == "insufficient_evidence"
    out_of_range = [c for c in claim_set.claims if c.period == "2021Q4"]
    assert out_of_range and result_by_cid[out_of_range[0].cid].status == "insufficient_evidence"


# ---------------- 核验报告导出 ----------------

def _small_report():
    frags = extract_fragments("r.md", "2024年营业收入为415亿元。\n2024年净利润为415亿元。".encode())
    snap = make_snap()
    return verify_claim_set(extract_claims(frags), snap,
                            fragments_text="", target_desc="单元测试报告")


def test_report_markdown_complete():
    md = report_to_markdown(_small_report())
    assert "核验报告" in md
    assert "通过" in md and "不通过" in md
    assert "人工复核清单" in md and "声明" in md
    counts_line = [l for l in md.splitlines() if "声明总数" in l][0]
    assert "通过 1" in counts_line and "不通过 1" in counts_line


def test_report_docx_complete():
    data = report_to_docx_bytes(_small_report())
    doc = Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("人工复核清单" in h for h in headings)
    assert doc.tables, "应有结果表格"
    rows = doc.tables[0].rows
    assert rows[0].cells[0].text == "ID"
    assert len(rows) >= 3  # 表头 + ≥2 条结果
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "独立" in body or "声明" in body
